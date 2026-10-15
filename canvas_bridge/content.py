"""Bounded text extraction. Page numbers are PDF pages or slide numbers, never guessed."""

from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import parse_qsl, urlencode, urljoin, urlsplit, urlunsplit
import xml.etree.ElementTree as ET
import zipfile

from .config import BridgeError


def clean_url(value: str, base: str = "", _depth: int = 0) -> str:
    if not value:
        return ""
    parts = urlsplit(urljoin(base, value))
    if parts.scheme not in ("http", "https") or parts.username or parts.password:
        return ""
    sensitive = {"access_token", "refresh_token", "id_token", "token", "auth", "authorization",
                 "code", "verifier", "signature", "sig", "credential", "jwt", "policy",
                 "key-pair-id", "awsaccesskeyid", "client_secret", "password", "api_key",
                 "apikey", "session_token", "samlresponse", "samlrequest"}

    def secret(key):
        lowered = key.lower().replace("-", "_")
        return key.lower() in sensitive or lowered in sensitive or lowered.startswith(("x_amz_", "x_goog_", "oauth_"))

    query = []
    for key, item in parse_qsl(parts.query, keep_blank_values=True):
        if secret(key):
            continue
        if item.startswith(("http://", "https://")):
            if _depth >= 2:
                continue
            item = clean_url(item, _depth=_depth + 1)
        query.append((key, item))
    fragment = parts.fragment
    if any(secret(k) for k, _ in parse_qsl(fragment)):
        fragment = ""
    return urlunsplit((parts.scheme, parts.netloc, parts.path, urlencode(query), fragment))


class HtmlText(HTMLParser):
    def __init__(self, base=""):
        super().__init__(convert_charrefs=True)
        self.base, self.chunks, self.links = base, [], []
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self.hidden += 1
        if not self.hidden:
            if tag in ("p", "div", "li", "br", "tr", "h1", "h2", "h3", "h4"):
                self.chunks.append("\n")
            attrs = dict(attrs)
            link = clean_url(attrs.get("href", ""), self.base) if tag == "a" else ""
            if link and link not in self.links:
                self.links.append(link)
            if tag == "img":
                self.chunks.append("[图片：" + attrs.get("alt", "无文字说明") + "]")

    def handle_endtag(self, tag):
        if tag in ("script", "style") and self.hidden:
            self.hidden -= 1
        if tag in ("p", "div", "li", "tr") and not self.hidden:
            self.chunks.append("\n")

    def handle_data(self, data):
        if not self.hidden:
            self.chunks.append(data)


def html_content(value: str | None, base=""):
    parser = HtmlText(base)
    parser.feed(value or "")
    lines = [line.strip() for line in "".join(parser.chunks).splitlines()]
    return {"text": "\n".join(line for line in lines if line), "links": parser.links}


def extract(path: Path, start: int = 1, count: int = 5, max_chars: int = 20000, offset: int = 0):
    if not 1 <= start or not 1 <= count <= 20 or not 100 <= max_chars <= 50000 or offset < 0:
        raise BridgeError("start 至少为 1，count 应为 1–20，max_chars 应为 100–50000，offset 至少为 0。")
    suffix = path.suffix.lower()
    warnings = []
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(path)
            if reader.is_encrypted and not reader.decrypt(""):
                raise BridgeError("PDF 已加密，无法提取文字。")
            total, unit = len(reader.pages), "pdf_page"
            selected = [(i + 1, reader.pages[i].extract_text() or "")
                        for i in range(start - 1, min(start - 1 + count, total))]
            warnings.append("仅提取文字；图片、扫描页、手写批注和公式可能缺失，需要视觉核对。")
        elif suffix == ".pptx":
            from pptx import Presentation
            deck = Presentation(path)
            total, unit = len(deck.slides), "slide"

            def shape_text(shape):
                if shape.has_text_frame:
                    yield shape.text
                if shape.has_table:
                    for row in shape.table.rows:
                        yield " | ".join(cell.text for cell in row.cells)
                if shape.shape_type == 6:  # GROUP
                    for child in shape.shapes:
                        yield from shape_text(child)

            selected = []
            for i in range(start - 1, min(start - 1 + count, total)):
                slide = deck.slides[i]
                texts = [t for shape in slide.shapes for t in shape_text(shape)]
                if slide.has_notes_slide:
                    notes = slide.notes_slide.notes_text_frame
                    if notes and notes.text.strip():
                        texts.append("[演讲者备注]\n" + notes.text)
                selected.append((i + 1, "\n".join(texts)))
            warnings.append("仅提取幻灯片文字、表格和备注；图形、公式和手写内容需要视觉核对。")
        elif suffix == ".docx":
            with zipfile.ZipFile(path) as archive:
                if archive.getinfo("word/document.xml").file_size > 20_000_000:
                    raise BridgeError("Word 正文过大，无法安全提取。")
                root = ET.fromstring(archive.read("word/document.xml"))
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            paragraphs = ["".join(n.text or "" for n in p.findall(".//w:t", ns))
                          for p in root.findall(".//w:p", ns)]
            total, unit = len(paragraphs), "paragraph"
            selected = list(enumerate(paragraphs[start - 1:start - 1 + count], start))
            warnings.append("Word 按段落定位，不代表实际页码；图像和批注未提取。")
        elif suffix in (".txt", ".md", ".csv", ".tsv", ".json", ".html", ".htm"):
            data = path.read_bytes()
            if len(data) > 10_000_000:
                raise BridgeError("文本文件过大，请下载后使用本地工具读取。")
            try:
                text = data.decode("utf-8-sig")
            except UnicodeDecodeError:
                try:
                    text = data.decode("gb18030")
                except UnicodeDecodeError:
                    text = data.decode("utf-8", errors="replace")
                    warnings.append("部分文字编码无法识别，已使用替代字符。")
            if suffix in (".html", ".htm"):
                text = html_content(text)["text"]
            lines = text.splitlines()
            total, unit = len(lines), "line"
            selected = list(enumerate(lines[start - 1:start - 1 + count], start))
        else:
            raise BridgeError("该文件格式暂不支持文字提取；可使用 download_material 下载原文件。")
    except BridgeError:
        raise
    except Exception:
        raise BridgeError("文件文字提取失败，原文件可能损坏或包含不支持的内容。") from None
    if start > total and total:
        raise BridgeError(f"start 超出范围，文件共有 {total} 个 {unit}。")
    if offset and (not selected or offset > len(selected[0][1])):
        raise BridgeError("offset 超出起始页或段落的文字范围，请使用返回的 next_cursor。")
    units, budget, next_start, next_offset, clipped = [], max_chars, None, 0, False
    for number, text in selected:
        if budget <= 0:
            next_start = number
            break
        begin = offset if number == start else 0
        shown = text[begin:begin + budget]
        is_clipped = begin + len(shown) < len(text)
        units.append({"number": number, "offset": begin, "text": shown, "truncated": is_clipped})
        budget -= len(shown)
        if not text.strip():
            warnings.append(f"{unit} {number} 未提取到文字，可能是扫描或图片内容。")
        if is_clipped:
            clipped = True
            next_start, next_offset = number, begin + len(shown)
            warnings.append(f"{unit} {number} 因字数限制被截断；请使用 next_cursor 继续读取同一页。")
            break
    if next_start is None and units and units[-1]["number"] < total:
        next_start = units[-1]["number"] + 1
    cursor = {"start": next_start, "offset": next_offset} if next_start is not None else None
    return {"unit": unit, "total_units": total, "units": units, "next_start": next_start,
            "next_offset": next_offset if cursor else None, "next_cursor": cursor,
            "text_truncated": clipped or len(units) < len(selected), "warnings": warnings}
