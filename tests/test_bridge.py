import io
from http.client import IncompleteRead
import json
from email.message import Message
from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock, patch
from urllib.error import HTTPError
import zipfile

from canvas_bridge.client import CanvasClient, valid_id
from canvas_bridge.config import BridgeError, Settings, read_token
from canvas_bridge.content import clean_url, extract, html_content
from canvas_bridge.service import CanvasService


TOKEN = "123~" + "a" * 64


def response(body, content_type="application/json", **headers):
    value = io.BytesIO(body if isinstance(body, bytes) else json.dumps(body).encode())
    value.headers = Message()
    value.headers["Content-Type"] = content_type
    for key, item in headers.items():
        value.headers[key] = item
    return value


class BridgeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.settings = Settings("https://canvas.example.edu", TOKEN, cache_dir=self.folder)
        self.client = CanvasClient(self.settings)

    def test_docx_split_runs_reloads_changed_token(self):
        file = self.folder / "key.docx"
        def write(token):
            with zipfile.ZipFile(file, "w") as archive:
                archive.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>' + token[:20] + '</w:t></w:r><w:r><w:t>' + token[20:] + '</w:t></w:r></w:p></w:body></w:document>')
        write(TOKEN)
        self.assertEqual(read_token(file), TOKEN)
        new = "456~" + "b" * 64
        write(new)
        self.assertEqual(read_token(file), new)
        self.assertNotIn(TOKEN, repr(self.settings))

    def test_ambiguous_token_document_rejected(self):
        file = self.folder / "key.docx"
        with zipfile.ZipFile(file, "w") as archive:
            archive.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:p><w:r><w:t>' + TOKEN + ' ' + 'b' * 64 + '</w:t></w:r></w:p></w:document>')
        with self.assertRaises(BridgeError):
            read_token(file)

    def test_pagination_follows_opaque_link_and_marks_truncation(self):
        self.client.opener.open = Mock(side_effect=[
            response([{"id": "1"}], link='<https://canvas.example.edu/api/v1/courses?cursor=x>; rel="next"'),
            response([{"id": "2"}, {"id": "3"}]),
        ])
        result = self.client.list("/api/v1/courses", limit=2)
        self.assertEqual([x["id"] for x in result["items"]], ["1", "2"])
        self.assertTrue(result["truncated"])
        self.assertEqual(self.client.opener.open.call_count, 2)

    def test_foreign_pagination_rejected_before_request(self):
        self.client.opener.open = Mock(return_value=response([], link='<https://evil.example/api/v1/courses>; rel="next"'))
        with self.assertRaises(BridgeError):
            self.client.list("/api/v1/courses")
        self.assertEqual(self.client.opener.open.call_count, 1)

    def test_chunk_transport_error_accepts_only_complete_json(self):
        broken = response(b"")
        broken.read = Mock(side_effect=IncompleteRead(b'[{"id":"1"}]'))
        self.client.opener.open = Mock(return_value=broken)
        result, _ = self.client.get("/api/v1/courses")
        self.assertEqual(result, [{"id": "1"}])

    def test_truncated_json_retries_instead_of_returning_partial(self):
        broken = response(b"")
        broken.read = Mock(side_effect=IncompleteRead(b'[{"id":'))
        self.client.opener.open = Mock(side_effect=[broken, response([{"id": "2"}])])
        with patch("canvas_bridge.client.time.sleep"):
            result, _ = self.client.get("/api/v1/courses")
        self.assertEqual(result, [{"id": "2"}])
        self.assertEqual(self.client.opener.open.call_count, 2)

    def test_truncated_file_is_not_accepted(self):
        self.client.opener.open = Mock(side_effect=[response(b"short", "application/pdf") for _ in range(3)])
        with patch("canvas_bridge.client.time.sleep"):
            with self.assertRaises(BridgeError):
                self.client.download("https://canvas.example.edu/files/1/download", expected_size=100)
        self.assertEqual(self.client.opener.open.call_count, 3)

    def test_api_redirect_never_followed(self):
        headers = Message()
        headers["Location"] = "https://evil.example/"
        self.client.opener.open = Mock(side_effect=HTTPError("hidden", 302, "redirect", headers, None))
        with self.assertRaises(BridgeError):
            self.client.get("/api/v1/courses")
        self.assertEqual(self.client.opener.open.call_count, 1)

    def test_download_redirect_strips_authorization(self):
        headers = Message()
        headers["Location"] = "https://cdn.example.edu/file.pdf?signature=hidden"
        requests = []
        def opened(request, timeout):
            requests.append(request)
            if len(requests) == 1:
                raise HTTPError("hidden", 302, "redirect", headers, None)
            return response(b"file contents", "application/pdf")
        self.client.opener.open = opened
        with patch("canvas_bridge.client.socket.getaddrinfo", return_value=[(2, 1, 6, "", ("8.8.8.8", 443))]):
            data = self.client.download("https://canvas.example.edu/files/1/download")
        self.assertEqual(data, b"file contents")
        self.assertEqual(requests[0].get_header("Authorization"), "Bearer " + TOKEN)
        self.assertIsNone(requests[1].get_header("Authorization"))
        self.assertTrue(all(request.get_method() == "GET" for request in requests))

    def test_private_download_rejected(self):
        self.client.opener.open = Mock()
        with patch("canvas_bridge.client.socket.getaddrinfo", return_value=[(2, 1, 6, "", ("127.0.0.1", 443))]):
            with self.assertRaises(BridgeError):
                self.client.download("https://cdn.example.edu/file")
        self.client.opener.open.assert_not_called()

    def test_fake_ip_proxy_exception_is_domain_and_range_scoped(self):
        self.settings.proxy_fake_ip_domains = ("canvas-user-content.com",)
        self.client.opener.open = Mock(return_value=response(b"data", "application/pdf"))
        with patch("canvas_bridge.client.socket.getaddrinfo", return_value=[(2, 1, 6, "", ("198.18.0.76", 443))]):
            self.assertEqual(self.client.download("https://a1-2.cluster42.canvas-user-content.com/file"), b"data")
            with self.assertRaises(BridgeError):
                self.client.download("https://canvas-user-content.com.evil.example/file")
        with patch("canvas_bridge.client.socket.getaddrinfo", return_value=[(2, 1, 6, "", ("127.0.0.1", 443))]):
            with self.assertRaises(BridgeError):
                self.client.download("https://a1-2.cluster42.canvas-user-content.com/file")

    def test_download_size_limit(self):
        self.client.opener.open = Mock(return_value=response(b"", "application/pdf", **{"Content-Length": str(51 * 1024 * 1024)}))
        with self.assertRaises(BridgeError):
            self.client.download("https://canvas.example.edu/files/1/download")

    def test_api_errors_do_not_echo_secrets(self):
        self.client.opener.open = Mock(side_effect=HTTPError(TOKEN, 401, TOKEN, Message(), None))
        with self.assertRaises(BridgeError) as error:
            self.client.get("/api/v1/courses")
        self.assertNotIn(TOKEN, str(error.exception))
        self.assertIn("401", str(error.exception))

    def test_html_links_keep_attachment_ids_without_signed_queries(self):
        value = html_content('<p>要求</p><a href="/courses/1/files/42?verifier=secret">附件</a><script>ignore me</script>', self.settings.base_url)
        self.assertIn("要求", value["text"])
        self.assertNotIn("ignore me", value["text"])
        self.assertEqual(value["links"], ["https://canvas.example.edu/courses/1/files/42"])

    def test_assignment_timezone_and_unknown_submission(self):
        service = CanvasService(self.settings, Mock())
        item = service._assignment({"id": "1", "due_at": "2026-10-15T15:59:00Z"})
        self.assertEqual(item["due_at_local"], "2026-10-15T23:59:00+08:00")
        self.assertIsNone(item["submission"])

    def test_files_permission_falls_back_to_accessible_module_links(self):
        client = Mock()
        client.list.side_effect = [BridgeError("HTTP 403：denied"),
            {"items": [{"id": "10", "name": "Week 1"}], "count": 1, "truncated": False},
            {"items": [{"id": "11", "type": "File", "content_id": "42"}], "count": 1, "truncated": False}]
        client.get.return_value = ({"id": "42", "filename": "lecture.pdf"}, {})
        result = CanvasService(self.settings, client).list_course_materials("1")
        self.assertEqual(result["items"][0]["id"], "42")
        self.assertTrue(result["partial"])
        self.assertEqual(result["discovery"], "module_links")
        self.assertTrue(result["warnings"])
        client.get.assert_not_called()

    def test_cache_refresh_and_redaction(self):
        client = Mock()
        metadata = {"id": "42", "filename": "../lecture.txt", "size": 6, "updated_at": "v1",
                    "url": "https://canvas.example.edu/files/42?verifier=secret"}
        client.get.return_value = (metadata, {})
        client.download.return_value = b"hello\n"
        service = CanvasService(self.settings, client)
        first = service.download_material("42")
        second = service.download_material("42")
        self.assertFalse(first["cached"])
        self.assertTrue(second["cached"])
        self.assertEqual(client.download.call_count, 1)
        self.assertEqual(Path(first["local_path"]).parent, self.folder / "files")
        metadata["updated_at"] = "v2"
        service.download_material("42")
        self.assertEqual(client.download.call_count, 2)
        self.assertNotIn("verifier", (self.folder / "files/42.json").read_text())
        self.assertNotIn(TOKEN, json.dumps(service._result("/api/v1/files", text=TOKEN)))

    def test_link_only_file_is_accessible_but_locked_file_is_blocked(self):
        client = Mock()
        metadata = {"id": "42", "filename": "homework.txt", "size": 2, "updated_at": "v1",
                    "url": "https://canvas.example.edu/files/42/download", "hidden_for_user": True,
                    "locked_for_user": False}
        client.get.return_value = (metadata, {})
        client.download.return_value = b"hi"
        service = CanvasService(self.settings, client)
        result = service.download_material("42")
        self.assertTrue(result["file"]["hidden_for_user"])
        client.download.assert_called_once()
        metadata["locked_for_user"] = True
        client.download.reset_mock()
        with self.assertRaises(BridgeError):
            service.download_material("42")
        client.download.assert_not_called()

    def test_resume_cursor_rejects_changed_file_version(self):
        path = self.folder / "long.txt"
        path.write_text("A" * 250)
        service = CanvasService(self.settings, Mock())
        service.download_material = Mock(return_value={"source_url": "https://canvas.example.edu/api/v1/files/42",
            "local_path": str(path), "sha256": "a" * 64, "file": {"id": "42"}})
        first = service.read_material("42", count=1, max_chars=100)
        self.assertEqual(first["next_cursor"]["expected_sha256"], "a" * 64)
        service.download_material.return_value["sha256"] = "b" * 64
        with self.assertRaises(BridgeError) as error:
            service.read_material("42", count=1, max_chars=100, **first["next_cursor"])
        self.assertIn("已更新", str(error.exception))

    def test_pdf_pages_and_missing_text_are_explicit(self):
        from pypdf import PdfWriter
        path = self.folder / "scanned.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.add_blank_page(width=100, height=100)
        writer.write(path)
        value = extract(path, start=2, count=1)
        self.assertEqual(value["total_units"], 2)
        self.assertEqual(value["units"][0]["number"], 2)
        self.assertTrue(any("未提取到文字" in x for x in value["warnings"]))

    def test_pptx_text_and_speaker_notes(self):
        from pptx import Presentation
        from pptx.util import Inches
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "Regression example"
        slide.notes_slide.notes_text_frame.text = "Lecturer clarification"
        path = self.folder / "lecture.pptx"
        deck.save(path)
        value = extract(path)
        self.assertEqual(value["unit"], "slide")
        self.assertIn("Regression example", value["units"][0]["text"])
        self.assertIn("Lecturer clarification", value["units"][0]["text"])

    def test_text_window_and_truncation(self):
        path = self.folder / "lecture.txt"
        path.write_text("first\n" + "x" * 300 + "\nthird", encoding="utf-8")
        value = extract(path, start=2, count=2, max_chars=100)
        self.assertTrue(value["text_truncated"])
        self.assertTrue(value["units"][0]["truncated"])
        self.assertEqual(value["next_start"], 2)
        self.assertEqual(value["next_cursor"], {"start": 2, "offset": 100})
        self.assertEqual(len(value["units"][0]["text"]), 100)

    def test_external_business_parameters_survive_secret_removal(self):
        value = clean_url("https://video.example/watch?v=lecture123&t=90&access_token=hidden&X-Amz-Signature=hidden#page=3")
        self.assertEqual(value, "https://video.example/watch?v=lecture123&t=90#page=3")
        self.assertEqual(clean_url("" , self.settings.base_url), "")
        nested = clean_url("https://tool.example/launch?id=3&return_to=https%3A%2F%2Fvideo.example%2Fwatch%3Fv%3Dabc%26token%3Dsecret")
        self.assertNotIn("secret", nested)
        self.assertIn("v%3Dabc", nested)
        self.assertNotIn("access_token", clean_url("https://example.edu/page#access_token=hidden"))

    def test_cursor_reassembles_long_units_without_skips_or_duplicates(self):
        path = self.folder / "lecture.txt"
        original = ["A" * 250 + "TAIL_ONE", "B" * 120 + "TAIL_TWO", "end"]
        path.write_text("\n".join(original), encoding="utf-8")
        cursor, pieces = {"start": 1, "offset": 0}, {1: "", 2: "", 3: ""}
        calls = 0
        while cursor is not None:
            value = extract(path, count=2, max_chars=100, **cursor)
            for unit in value["units"]:
                self.assertEqual(unit["offset"], len(pieces[unit["number"]]))
                pieces[unit["number"]] += unit["text"]
            cursor = value["next_cursor"]
            calls += 1
            self.assertLess(calls, 10)
        self.assertEqual(list(pieces.values()), original)

    def test_invalid_cursor_is_rejected(self):
        path = self.folder / "short.txt"
        path.write_text("hello")
        with self.assertRaises(BridgeError):
            extract(path, offset=100)

    def test_invalid_ids_block_path_injection(self):
        for value in ["../1", "0", "https://evil.example", "1?access_token=x"]:
            with self.assertRaises(BridgeError):
                valid_id(value)


if __name__ == "__main__":
    unittest.main()
